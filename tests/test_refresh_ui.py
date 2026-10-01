"""Prueba de sincronización sin escribir en la Google Sheets del usuario."""
import os
from pathlib import Path
import socket
import subprocess
import sys
import time

from openpyxl import Workbook
from playwright.sync_api import sync_playwright, expect
import pytest
import requests


def write_source(path, rows, target=4):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'registros'
    sheet.append(['id_registro', 'fecha_inicio', 'nombre_proyecto', 'recursos_gestionados',
                  'recursos_asignados', 'recursos_institucionales'])
    for identifier, name in rows:
        sheet.append([identifier, '2026-02-09', name, 100, 50, 0])
    definitions = workbook.create_sheet('indicadores')
    definitions.append(['id_indicador', 'nombre', 'periodicidad', 'modo', 'numerador', 'meta'])
    definitions.append(['I1', 'Gestiones de prueba', 'ANUAL', 'META', '{"op":"CONTAR"}', target])
    closures = workbook.create_sheet('cierres')
    closures.append(['id_cierre', 'id_indicador', 'anio', 'periodo', 'valor', 'meta', 'cumplimiento', 'banda'])
    closures.append(['C1', 'I1', 2026, 'ANUAL', 8, 10, 0.8, 'RIESGO'])
    temporary = path.with_suffix('.next.xlsx')
    workbook.save(temporary)
    temporary.replace(path)


def test_two_sessions_observe_updates_and_recover_from_invalid_source(request, tmp_path):
    if not request.config.getoption('--ui'):
        pytest.skip('La prueba de dos sesiones requiere --ui.')
    source = tmp_path / 'test-source.xlsx'
    write_source(source, [('R1', 'Inicial'), ('R0', 'Eliminar')])
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
    environment = {**os.environ, 'SOURCE_MODE': 'file', 'LOCAL_WORKBOOK': str(source),
                   'REFRESH_SECONDS': '5', 'GEMINI_API_KEY': ''}
    app_dir = Path(__file__).resolve().parents[1]
    with (tmp_path / 'server.log').open('w', encoding='utf-8') as log:
        process = subprocess.Popen(
            [sys.executable, '-m', 'shiny', 'run', '--host', '127.0.0.1', '--port', str(port), str(app_dir / 'app.py')],
            env=environment, stdout=log, stderr=log,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0,
        )
        url = f'http://127.0.0.1:{port}'
        try:
            deadline = time.monotonic() + 30
            while True:
                try:
                    if requests.get(url, timeout=1).status_code == 200:
                        break
                except requests.RequestException:
                    pass
                if time.monotonic() > deadline or process.poll() is not None:
                    pytest.fail('No inició la app de prueba: ' + (tmp_path / 'server.log').read_text(encoding='utf-8'))
                time.sleep(0.2)
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                pages = [browser.new_page(), browser.new_page()]
                errors = []
                for page in pages:
                    page.on('pageerror', lambda error: errors.append(str(error)))
                    page.goto(url)
                    expect(page.locator('#source_status')).to_contain_text('Archivo local', timeout=30000)
                    page.locator('#year').select_option('2026')
                    expect(page.locator('.metric-value').first).to_have_text('2', timeout=10000)
                    page.locator('.sidebar label:has(input[value="indicadores"])').click()
                    expect(page.locator('.indicator-card')).to_contain_text('50,0 %', timeout=10000)
                    page.locator('.sidebar label:has(input[value="analisis"])').click()
                # Alta, cambio y baja en un mismo ciclo; ambas sesiones se actualizan solas.
                write_source(source, [('R1', 'Actualizado'), ('R2', 'Nuevo'), ('R3', 'Tercero')], target=3)
                for page in pages:
                    expect(page.locator('.metric-value').first).to_have_text('3', timeout=15000)
                    page.locator('.sidebar label:has(input[value="indicadores"])').click()
                    expect(page.locator('.indicator-card')).to_contain_text('100,0 %', timeout=10000)
                    page.locator('.sidebar label:has(input[value="cierres"])').click()
                    expect(page.locator('#main_table')).to_contain_text('80', timeout=10000)
                    expect(page.locator('#main_table')).to_contain_text('En riesgo')
                    page.locator('.sidebar label:has(input[value="registros"])').click()
                    expect(page.locator('#record_id')).to_contain_text('Actualizado', timeout=10000)
                    expect(page.locator('#record_id')).to_contain_text('Nuevo')
                    expect(page.locator('#record_id')).not_to_contain_text('Inicial')
                    expect(page.locator('#record_id')).not_to_contain_text('Eliminar')
                # Una lectura inválida conserva los últimos datos, con aviso visible.
                source.write_bytes(b'fuente de prueba invalida')
                for page in pages:
                    expect(page.locator('#source_status')).to_contain_text('No se pudo actualizar', timeout=15000)
                    expect(page.locator('#record_id')).to_contain_text('Actualizado')
                write_source(source, [('R4', 'Recuperado')])
                for page in pages:
                    expect(page.locator('#source_status')).to_contain_text('Conectado', timeout=15000)
                    expect(page.locator('#record_id')).to_contain_text('Recuperado', timeout=10000)
                    expect(page.locator('#record_id')).not_to_contain_text('Actualizado')
                assert not errors, errors
                browser.close()
        finally:
            process.terminate()
            process.wait(timeout=15)
