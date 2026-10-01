from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from indicadores.normalization import DataValidationError, normalize_tables, parse_date, parse_number, read_workbook


@pytest.mark.parametrize('value,expected', [
    ('1.234.567,50', 1234567.5), ('$ 1.728.000', 1728000),
    ('0,8', 0.8), ('80%', 0.8), ('3.34E8', 334000000),
    ('1,234,567.50', 1234567.5), (0, 0), ('', None), (None, None),
])
def test_numbers_keep_missing_separate_from_zero(value, expected):
    assert parse_number(value) == expected


@pytest.mark.parametrize('value', ['sin dato', float('nan'), '1.2.3,4', 'infinity'])
def test_invalid_amount_is_not_silently_zero(value):
    with pytest.raises(ValueError):
        parse_number(value)


@pytest.mark.parametrize('value,expected', [
    ('2026-02-09', date(2026, 2, 9)), ('09/02/2026', date(2026, 2, 9)),
    (datetime(2026, 2, 9, 13, 5), date(2026, 2, 9)),
    (46287, date(2026, 9, 22)), ('', None),
])
def test_dates_support_sheet_and_excel_formats(value, expected):
    assert parse_date(value) == expected


def test_duplicate_ids_reject_ambiguous_snapshot():
    with pytest.raises(DataValidationError, match='duplicado'):
        normalize_tables({'registros': [
            {'id_registro': 'R1', 'fecha_inicio': '2026-02-09'},
            {'id_registro': 'R1', 'fecha_inicio': '2026-02-10'},
        ]})


def test_bad_fields_have_visible_issues_and_preserve_valid_rows():
    snapshot = normalize_tables({'registros': [
        {'id_registro': 'R1', 'fecha_inicio': '2026-02-09', 'recursos_asignados': 'incorrecto'},
        {'id_registro': 'R2', 'fecha_inicio': 'fecha inválida', 'recursos_asignados': 0},
    ]})
    assert len(snapshot.tables['registros']) == 2
    assert len(snapshot.issues) == 2
    assert snapshot.tables['registros'][0]['recursos_asignados'] is None
    assert snapshot.tables['registros'][1]['recursos_asignados'] == 0


def test_workbook_reads_canonical_tables_only_and_keeps_codes_as_text():
    wb = Workbook()
    ws = wb.active
    ws.title = 'objetivos_corporativos'
    ws.append(['codigo', 'nombre'])
    ws.append(['4.10', 'Objetivo'])
    ws2 = wb.create_sheet('registros')
    ws2.append(['id_registro', 'fecha_inicio', 'nombre_proyecto'])
    ws2.append(['R1', datetime(2026, 2, 9), 'Proyecto'])
    old = wb.create_sheet('copia para propuesta de variabl')
    old.append(['#', 'NOMBRE'])
    old.append([1, 'Proyecto'])
    buffer = BytesIO()
    wb.save(buffer)
    snapshot = read_workbook(buffer.getvalue())
    assert len(snapshot.tables['registros']) == 1
    assert snapshot.tables['objetivos_corporativos'][0]['codigo'] == '4.10'
    assert 'copia para propuesta de variabl' not in snapshot.tables


def test_missing_primary_key_is_not_manufactured():
    with pytest.raises(DataValidationError, match='identificador'):
        normalize_tables({'registros': [{'nombre_proyecto': 'Gestión', 'fecha_inicio': '2026-02-09'}]})


def test_catalog_values_are_text_even_when_column_is_named_valor():
    snapshot = normalize_tables({'catalogos': [{'categoria': 'proceso', 'valor': 'GESTIÓN DE PROYECTOS', 'orden': 1}]})
    assert snapshot.rows('catalogos')[0]['valor'] == 'GESTIÓN DE PROYECTOS'
    assert not snapshot.issues
